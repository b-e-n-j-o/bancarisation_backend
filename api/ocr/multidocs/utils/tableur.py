"""
tableur.py — XLSX → blocs "table logique".

C'est le normaliseur le plus riche, parce que c'est le format le plus hostile :
un classeur de BE empile plusieurs tables dans une même feuille, fusionne des
cellules de montants sur plusieurs lignes, et cache ses sous-totaux dans des
formules.

Grain retenu : UNE TABLE LOGIQUE PAR BLOC (bande de lignes non vides), locator
"Feuille!debut:fin". Ce grain sert deux fois :
  · ancre vérifiable ("feuille Planning, lignes 8 à 40") ;
  · unité de fan-out — on envoie une table au LLM, pas un classeur entier.
    (Ton premier test d'extraction budget a explosé le quota de réflexion sur
    le classeur entier : le modèle tournait en rond sur des contradictions
    inter-tables. Une table à la fois supprime le problème à la racine.)

Annotations conservées, reprises telles quelles par les prompts d'extraction :
  ⟨montants communs aux lignes X-Y⟩   fusion verticale, ligne d'ancrage
  ⟨montants portés par la ligne X⟩    fusion verticale, lignes suiveuses
  ⟨col. A-D⟩                          fusion horizontale
  [TOTAL?]                            la ligne contient une formule de somme
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Optional

from .corpus import Bloc, DocumentNormalise, TypeBloc

MAX_COLONNES = 40
MAX_LIGNES_TABLE = 400  # au-delà, on découpe : une table doit rester "prompt-able"


def normaliser_xlsx(chemin: Path, doc_id: str, **_) -> DocumentNormalise:
    from openpyxl import load_workbook

    # Deux passes obligatoires : data_only=True donne les valeurs sans les
    # formules, le défaut donne les formules sans les valeurs.
    wb_val = load_workbook(chemin, data_only=True, read_only=False)
    wb_for = load_workbook(chemin, data_only=False, read_only=False)

    blocs: list[Bloc] = []
    for ws in wb_val.worksheets:
        wsf = wb_for[ws.title]
        fusions = _index_fusions(ws)
        for debut, fin in _bandes(ws):
            texte = _table_markdown(ws, wsf, debut, fin, fusions)
            if not texte.strip():
                continue
            blocs.append(
                Bloc(
                    locator=f"{ws.title}!{debut}:{fin}",
                    type=TypeBloc.tableau,
                    texte=texte,
                    titre=ws.title,
                    meta={
                        "feuille": ws.title,
                        "lignes": [debut, fin],
                        "nb_lignes": fin - debut + 1,
                    },
                )
            )

    return DocumentNormalise(
        doc_id=doc_id,
        nom_fichier=chemin.name,
        format="xlsx",
        sha256="",
        blocs=blocs,
        meta={"feuilles": [ws.title for ws in wb_val.worksheets]},
    )


# --- découpage en tables logiques ------------------------------------------


def _ligne_vide(ws, ligne: int, max_col: int) -> bool:
    for col in range(1, max_col + 1):
        v = ws.cell(row=ligne, column=col).value
        if v is not None and str(v).strip() != "":
            return False
    return True


def _bandes(ws) -> list[tuple[int, int]]:
    """Bandes de lignes non vides séparées par au moins une ligne vide.
    Heuristique volontairement bête : elle sur-découpe plutôt que de fusionner
    deux tables distinctes — une sur-découpe se recolle au triage, une fusion
    abusive contamine l'extraction."""
    max_col = min(ws.max_column or 1, MAX_COLONNES)
    max_row = ws.max_row or 0
    bandes: list[tuple[int, int]] = []
    debut: Optional[int] = None

    for ligne in range(1, max_row + 1):
        vide = _ligne_vide(ws, ligne, max_col)
        if not vide and debut is None:
            debut = ligne
        elif vide and debut is not None:
            bandes.extend(_decouper(debut, ligne - 1))
            debut = None
    if debut is not None:
        bandes.extend(_decouper(debut, max_row))

    return [(d, f) for d, f in bandes if f >= d]


def _decouper(debut: int, fin: int) -> list[tuple[int, int]]:
    if fin - debut + 1 <= MAX_LIGNES_TABLE:
        return [(debut, fin)]
    morceaux = []
    curseur = debut
    while curseur <= fin:
        borne = min(curseur + MAX_LIGNES_TABLE - 1, fin)
        morceaux.append((curseur, borne))
        curseur = borne + 1
    return morceaux


# --- fusions ---------------------------------------------------------------


def _index_fusions(ws) -> dict[tuple[int, int], dict]:
    """(ligne, colonne) → description de la fusion qui la couvre."""
    index: dict[tuple[int, int], dict] = {}
    for plage in ws.merged_cells.ranges:
        info = {
            "min_row": plage.min_row,
            "max_row": plage.max_row,
            "min_col": plage.min_col,
            "max_col": plage.max_col,
        }
        for r in range(plage.min_row, plage.max_row + 1):
            for c in range(plage.min_col, plage.max_col + 1):
                index[(r, c)] = info
    return index


def _annotations_ligne(ligne: int, max_col: int, fusions: dict) -> list[str]:
    notes: list[str] = []
    vues: set[tuple] = set()
    for col in range(1, max_col + 1):
        info = fusions.get((ligne, col))
        if not info:
            continue
        cle = (info["min_row"], info["max_row"], info["min_col"], info["max_col"])
        if cle in vues:
            continue
        vues.add(cle)

        if info["max_row"] > info["min_row"]:
            if ligne == info["min_row"]:
                notes.append(
                    f"⟨montants communs aux lignes {info['min_row']}-{info['max_row']}⟩"
                )
            else:
                notes.append(f"⟨montants portés par la ligne {info['min_row']}⟩")
        if info["max_col"] > info["min_col"]:
            notes.append(f"⟨col. {_lettre(info['min_col'])}-{_lettre(info['max_col'])}⟩")
    return notes


# --- rendu -----------------------------------------------------------------


def _lettre(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def _cellule(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, (dt.datetime, dt.date)):
        return valeur.isoformat()[:10]
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).replace("\n", " ").replace("|", "/").strip()


def _est_somme(wsf, ligne: int, max_col: int) -> bool:
    for col in range(1, max_col + 1):
        v = wsf.cell(row=ligne, column=col).value
        if isinstance(v, str) and v.startswith("=") and "SUM(" in v.upper():
            return True
    return False


def _table_markdown(ws, wsf, debut: int, fin: int, fusions: dict) -> str:
    max_col = min(ws.max_column or 1, MAX_COLONNES)
    # On ne rend que les colonnes réellement utilisées dans la bande.
    colonnes = [
        c for c in range(1, max_col + 1)
        if any(
            ws.cell(row=r, column=c).value not in (None, "")
            for r in range(debut, fin + 1)
        )
    ]
    if not colonnes:
        return ""

    entete = ["ligne"] + [_lettre(c) for c in colonnes]
    lignes_md = [
        "| " + " | ".join(entete) + " |",
        "| " + " | ".join("---" for _ in entete) + " |",
    ]

    for r in range(debut, fin + 1):
        cellules = [_cellule(ws.cell(row=r, column=c).value) for c in colonnes]
        etiquettes = _annotations_ligne(r, max_col, fusions)
        if _est_somme(wsf, r, max_col):
            etiquettes.append("[TOTAL?]")
        libelle = str(r) + (" " + " ".join(etiquettes) if etiquettes else "")
        lignes_md.append("| " + " | ".join([libelle] + cellules) + " |")

    return "\n".join(lignes_md)