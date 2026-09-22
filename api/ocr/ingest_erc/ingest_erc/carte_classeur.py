"""Carte du classeur Excel — passe 1.

Détection déterministe des tableaux, puis un appel LLM (rôles, colonnes, taxe)
vérifié et mis en cache par sha256. Les faits (conventions HT/TTC, codes rencontrés,
rôles) alimentent la réconciliation avant verrouillage.

La passe 3 ne fait plus cet appel : elle relit le cache et applique les rôles validés.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal, Optional

from openpyxl.utils import get_column_letter
from pydantic import Field, field_validator

from . import llm
from .carte import _IgnoreNull, _literal, Regles
from .excel_profil import _annee, _rang, code_norm, lire_grille
from .inventaire import norm
from .modeles import Fait, Source

RoleTableau = Literal["planning_passages", "planning_montants", "couts_unitaires", "decompte_realise",
                      "statuts", "synthese", "parcelles", "autre"]
Orientation = Literal["annees_en_colonnes", "annees_en_lignes", "sans_temps"]
ROLES_BUDGET = {"planning_montants", "couts_unitaires", "decompte_realise"}
CHAMPS_COL = ("action", "ug", "annee", "montant", "quantite", "prix_unitaire", "prestataire", "statut")


@dataclass
class TableauBrut:
    """Région dense d'une feuille avec sa ligne d'en-tête, détectée sans a priori métier."""
    id: str
    doc: str
    feuille: str
    titre: str
    lignes_entete: list[int]
    lignes: list[int]
    entetes: dict = field(default_factory=dict)     # col -> libellé (lignes d'en-tête jointes)


class TableauCarte(_IgnoreNull):
    id: str
    role: RoleTableau = "autre"
    orientation: Orientation = "sans_temps"
    col_action: Optional[str] = Field(None, description="libellé EXACT de l'en-tête qui identifie l'action (code ou libellé)")
    col_ug: Optional[str] = None
    col_annee: Optional[str] = None
    col_montant: Optional[str] = None
    col_quantite: Optional[str] = None
    col_prix_unitaire: Optional[str] = None
    col_prestataire: Optional[str] = None
    col_statut: Optional[str] = None
    taxe: Literal["HT", "TTC", "inconnu"] = "inconnu"

    @field_validator("role", mode="before")
    @classmethod
    def _v_role(cls, v):
        return _literal(RoleTableau, "autre")(v)

    @field_validator("orientation", mode="before")
    @classmethod
    def _v_orientation(cls, v):
        return _literal(Orientation, "sans_temps")(v)

    @field_validator("taxe", mode="before")
    @classmethod
    def _v_taxe(cls, v):
        return _literal(Literal["HT", "TTC", "inconnu"], "inconnu")(v)


class CarteClasseur(_IgnoreNull):
    tableaux: list[TableauCarte] = Field(default_factory=list)
    remarques: list[str] = Field(default_factory=list)
    verification: dict = Field(default_factory=dict)


PROMPT_CLASSEUR = """Tu es l'analyste documentaire d'une plateforme de suivi des mesures
compensatoires. Tu reçois la description des tableaux détectés dans un classeur Excel d'un
bureau d'études (en-têtes, colonnes temporelles résumées, 3 lignes d'exemple).
Tu ne lis pas le contenu métier : tu dis COMMENT chaque tableau est fait.

Pour chaque tableau [id] :
- role : planning_passages (quand une action a lieu : croix, codes ou nombres de passages),
  planning_montants (montants par année), couts_unitaires (prix par prestation sans axe
  temporel), decompte_realise (facturé / réalisé), statuts (réalisé / projeté / supprimé),
  synthese (lien action ↔ unité), parcelles, autre.
- orientation : annees_en_colonnes, annees_en_lignes, sans_temps.
- col_* : le libellé EXACT de l'en-tête (copié tel qu'affiché) de la colonne qui porte
  l'action (code ou libellé), l'unité de gestion, l'année (si en lignes), le montant, la
  quantité, le prix unitaire, le prestataire, le statut. Vide si la colonne n'existe pas.
- taxe : HT ou TTC seulement si c'est écrit dans le tableau, sinon inconnu.
Les libellés seront vérifiés : tout libellé introuvable est ignoré. N'invente rien.
Réponds uniquement avec un JSON conforme au schéma fourni."""


def _entete_like(cel) -> bool:
    return cel.typ == "txt" or _annee(cel.v) is not None


def _parts(lib: str) -> list[str]:
    return [p.strip() for p in str(lib).split(" / ")]


def _col_temps(lib: str) -> tuple[Optional[int], Optional[int]]:
    """(année, rang) lus dans un libellé d'en-tête « 2023 / N+1 »."""
    an = rg = None
    for p in _parts(lib):
        an = an if an is not None else _annee(p)
        rg = rg if rg is not None else _rang(p)
    return an, rg


def detecter_tableaux(feuilles, nom_doc: str) -> list[TableauBrut]:
    out = []
    for f in feuilles:
        occ = {}
        for r in range(1, f.nrow + 1):
            cs = [c for c in range(1, f.ncol + 1) if f.get(r, c).typ != "vide"]
            if len(cs) >= 2:
                occ[r] = cs
        r, k = 1, 0
        while r <= f.nrow:
            if r not in occ:
                r += 1
                continue
            bande = []
            while r in occ:
                bande.append(r)
                r += 1
            ratio = lambda rr: sum(_entete_like(f.get(rr, c)) for c in occ[rr]) / len(occ[rr])
            i = next((j for j, rr in enumerate(bande) if ratio(rr) >= 0.6), None)
            if i is None or i >= len(bande) - 1:
                continue
            entete = [bande[i]]
            for rr in bande[i + 1:i + 3]:
                rangs = sum(_rang(f.get(rr, c).v) is not None for c in occ[rr])
                if rangs >= max(2, len(occ[rr]) // 2):
                    entete.append(rr)
                else:
                    break
            donnees = bande[i + len(entete):]
            if not donnees:
                continue
            cols = sorted({c for rr in bande for c in occ[rr]})
            libs = {}
            for c in cols:
                parts = []
                for h in entete:
                    v = f.get(h, c).v
                    if v is not None and (not parts or parts[-1] != str(v)):
                        parts.append(str(v).replace(".0", "") if isinstance(v, float) else str(v))
                if parts:
                    libs[c] = " / ".join(parts)
            titre = ""
            for up in range(bande[0] - 1, max(0, bande[0] - 4), -1):
                txt = [x.v for x in f.ligne(up) if x.typ == "txt" and x.ancre[0] == up]
                if txt:
                    titre = str(txt[0])
                    break
            k += 1
            out.append(TableauBrut(f"{f.nom}#{k}", nom_doc, f.nom, titre, entete, donnees, libs))
    return out


def decrire_tableaux(tabs: list[TableauBrut], feuilles) -> str:
    """Description compacte envoyée au LLM : en-têtes (années résumées) + 3 lignes d'exemple."""
    par_nom = {f.nom: f for f in feuilles}
    blocs = []
    for t in tabs:
        f = par_nom[t.feuille]
        temps = [c for c, l in t.entetes.items() if _col_temps(l) != (None, None)]
        autres = [c for c in t.entetes if c not in temps]
        lignes = [f"[{t.id}] feuille « {t.feuille} » · titre « {t.titre} » · {len(t.lignes)} lignes"]
        lignes.append("  en-têtes : " + " | ".join(f"{get_column_letter(c)}={t.entetes[c]}" for c in autres))
        if temps:
            a0, r0 = _col_temps(t.entetes[temps[0]])
            a1, r1 = _col_temps(t.entetes[temps[-1]])
            lignes.append(f"  colonnes temporelles {get_column_letter(temps[0])}…{get_column_letter(temps[-1])} : "
                          f"{a0 if a0 is not None else 'rang ' + str(r0)} → {a1 if a1 is not None else 'rang ' + str(r1)}"
                          f" ({len(temps)} colonnes)")
        for r in t.lignes[:3]:
            vals = [f"{get_column_letter(c)}={str(f.get(r, c).v)[:30]}" for c in autres if f.get(r, c).typ != "vide"]
            vals += [f"{get_column_letter(c)}={str(f.get(r, c).v)[:12]}" for c in temps[:4] if f.get(r, c).typ != "vide"]
            lignes.append("  ex. : " + " | ".join(vals))
        blocs.append("\n".join(lignes))
    return "\n\n".join(blocs)


def _resoudre_col(lib: Optional[str], t: TableauBrut) -> Optional[int]:
    if not lib:
        return None
    n = norm(lib).strip()
    for c, l in t.entetes.items():
        if norm(l).strip() == n:
            return c
    for c, l in t.entetes.items():
        if n and (n in norm(l) or norm(l).strip() in n) and len(norm(l).strip()) >= 3:
            return c
    return None


def verifier_classeur(carte: CarteClasseur, tabs: list[TableauBrut]) -> tuple[CarteClasseur, dict]:
    """Écarte les tableaux inconnus et les libellés introuvables. Renvoie aussi, par tableau,
    l'index de colonne résolu pour chaque champ."""
    par_id = {t.id: t for t in tabs}
    rapport = {"tableaux": [0, len(carte.tableaux)], "colonnes": [0, 0], "ecartes": []}
    gardes, cols = [], {}
    for tc in carte.tableaux:
        t = par_id.get(tc.id)
        if t is None:
            rapport["ecartes"].append(f"tableau {tc.id} inconnu")
            continue
        res = {}
        for ch in CHAMPS_COL:
            lib = getattr(tc, f"col_{ch}")
            if not lib:
                continue
            rapport["colonnes"][1] += 1
            c = _resoudre_col(lib, t)
            if c is None:
                rapport["ecartes"].append(f"{tc.id} : colonne « {lib} » introuvable")
                setattr(tc, f"col_{ch}", None)
            else:
                res[ch] = c
                rapport["colonnes"][0] += 1
        if tc.taxe != "inconnu":
            visible = norm(" ".join([t.titre, t.feuille, *t.entetes.values()]))
            if not re.search(rf"\b{tc.taxe.lower()}\b", visible):
                rapport["ecartes"].append(f"{tc.id} : taxe {tc.taxe} non écrite dans le tableau")
                tc.taxe = "inconnu"
        gardes.append(tc)
        cols[tc.id] = res
        rapport["tableaux"][0] += 1
    carte.tableaux = gardes
    carte.verification = rapport
    return carte, cols


def carte_repli(tabs: list[TableauBrut], feuilles, regles: Optional[Regles] = None) -> CarteClasseur:
    """Sans LLM : seuls les tableaux à colonnes d'années ET à colonne de codes connus sont lus."""
    par_nom = {f.nom: f for f in feuilles}
    out = []
    for t in tabs:
        f = par_nom[t.feuille]
        temps = [c for c, l in t.entetes.items() if _col_temps(l) != (None, None)]
        if len(temps) < 3:
            continue
        def est_code(v):
            return (regles.code(v) if regles else code_norm(v)) is not None
        col_code = next((c for c in t.entetes if c not in temps
                         and sum(est_code(f.get(r, c).v) for r in t.lignes) >= len(t.lignes) / 2), None)
        if col_code is None:
            continue
        num = sum(f.get(r, c).typ == "num" for r in t.lignes for c in temps)
        txt = sum(f.get(r, c).typ == "txt" for r in t.lignes for c in temps)
        out.append(TableauCarte(id=t.id, role="planning_montants" if num > txt else "planning_passages",
                                orientation="annees_en_colonnes", col_action=t.entetes[col_code]))
    return CarteClasseur(tableaux=out, remarques=["carte de repli (sans LLM)"])


def _role() -> dict:
    from . import config
    return {
        "modele": config.modele_pour("carte_classeur"),
        "effort": config.effort_pour("carte_classeur"),
        "max_tokens": config.max_tokens_pour("carte_classeur"),
    }


def chemin_cache(doc, cache_dir: Optional[Path]) -> Optional[Path]:
    return cache_dir / f"{doc.sha256}.classeur.json" if cache_dir else None


def charger_ou_cartographier_classeur(doc, feuilles, tabs, cache_dir, regles=None
                                     ) -> tuple[CarteClasseur, dict]:
    """Passe 1 : cache sha256, sinon un appel LLM, sinon repli déterministe."""
    cache = chemin_cache(doc, cache_dir)
    if cache and cache.exists():
        carte = CarteClasseur.model_validate_json(cache.read_text())
    elif llm.actif():
        schema = json.dumps(CarteClasseur.model_json_schema(), ensure_ascii=False)
        carte = llm._appel(PROMPT_CLASSEUR + "\n\nSchéma JSON :\n" + schema,
                           decrire_tableaux(tabs, feuilles), CarteClasseur, **_role())
        if cache:
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_text(carte.model_dump_json(indent=1))
        print(f"🗺️  carte classeur {doc.nom} : {len(carte.tableaux)} tableau(x)", flush=True)
    else:
        carte = carte_repli(tabs, feuilles, regles)
    return verifier_classeur(carte, tabs)


def charger_carte_cache(doc, cache_dir) -> Optional[CarteClasseur]:
    """Passe 3 : relit le cache, sans nouvel appel LLM."""
    cache = chemin_cache(doc, cache_dir)
    if cache and cache.exists():
        return CarteClasseur.model_validate_json(cache.read_text())
    return None


def vue_classeur(doc, carte: CarteClasseur, tabs: list[TableauBrut]) -> dict:
    par_id = {t.id: t for t in tabs}
    return {
        "doc": doc.nom,
        "verification": carte.verification,
        "remarques": carte.remarques,
        "tableaux": [
            {
                "id": tc.id,
                "role": tc.role,
                "taxe": tc.taxe,
                "orientation": tc.orientation,
                "col_action": tc.col_action,
                "titre": (par_id[tc.id].titre if tc.id in par_id else ""),
                "feuille": (par_id[tc.id].feuille if tc.id in par_id else ""),
                "nb_lignes": len(par_id[tc.id].lignes) if tc.id in par_id else 0,
            }
            for tc in carte.tableaux
        ],
    }


def faits_depuis_carte(doc, carte: CarteClasseur, tabs: list[TableauBrut],
                       cols: dict, feuilles, regles=None) -> list[Fait]:
    """Conventions HT/TTC, rôles, codes rencontrés dans les colonnes action."""
    def _f(type_, cle, valeur, loc, extrait=None):
        return Fait(type=type_, cle=cle, valeur=valeur, methode="llm",
                    source=Source(doc=doc.nom, loc=loc, extrait=(extrait or "")[:160] or None))

    par_nom = {f.nom: f for f in feuilles}
    par_id = {t.id: t for t in tabs}
    out: list[Fait] = []
    for tc in carte.tableaux:
        t = par_id.get(tc.id)
        loc = tc.id
        extra = f"{t.titre} ({t.feuille})" if t else tc.id
        out.append(_f("tableau_role", tc.id,
                      {"role": tc.role, "taxe": tc.taxe, "orientation": tc.orientation,
                       "titre": t.titre if t else "", "feuille": t.feuille if t else "",
                       "doc": doc.nom},
                      loc, extra))
        if tc.taxe != "inconnu" and tc.role in ROLES_BUDGET:
            out.append(_f("convention_taxe", doc.nom, tc.taxe, loc,
                          f"{extra} : montants {tc.taxe}"))

        c_act = (cols.get(tc.id) or {}).get("action")
        if t is None or c_act is None:
            continue
        f = par_nom.get(t.feuille)
        if f is None:
            continue
        vus = set()
        for r in t.lignes:
            brut = f.get(r, c_act).v
            if brut in (None, ""):
                continue
            code = (regles.code(brut) if regles else code_norm(brut))
            if not code or code in vus:
                continue
            vus.add(code)
            out.append(_f("code_tableur", code, {"libelle": str(brut).strip(), "tableau": tc.id},
                          f"{t.feuille}!{get_column_letter(c_act)}{r}", str(brut).strip()))
    return out


def profiler_classeur(doc, cache_dir, regles=None) -> tuple[CarteClasseur, list[TableauBrut], dict, list[Fait], dict]:
    """Point d'entrée passe 1 : détection + carte + faits + vue."""
    feuilles = lire_grille(doc.chemin)
    tabs = detecter_tableaux(feuilles, doc.nom)
    carte, cols = charger_ou_cartographier_classeur(doc, feuilles, tabs, cache_dir, regles)
    faits = faits_depuis_carte(doc, carte, tabs, cols, feuilles, regles)
    return carte, tabs, cols, faits, vue_classeur(doc, carte, tabs)
