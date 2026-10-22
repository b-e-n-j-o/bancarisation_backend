"""Lecture de projet.xlsx. Les en-têtes sont reconnus par leur libellé."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from api.reimport.modele.constats import Constat
from api.reimport.modele.listes import DISPOSITIFS, ROLES, VERSION_MODELE
from api.reimport.modele.normaliser import (
    annee,
    code,
    code_ug,
    date_fr,
    liste_codes,
    mois,
    nombre,
    recurrence,
    statut,
    texte,
    type_ug,
)

ONGLETS = ("Projet", "Mesures", "UG", "Actions – plan", "Actions – calendrier", "Contacts")


def lire_classeur(chemin: Path) -> tuple[dict, list[Constat]]:
    if chemin.suffix.lower() == ".xlsm":
        return {}, [Constat("erreur", None, None, None, "Les fichiers .xlsm ne sont pas acceptés.")]
    wb = load_workbook(chemin, read_only=True, data_only=True)
    constats: list[Constat] = []
    if "_meta" not in wb.sheetnames:
        return {}, [Constat("erreur", None, None, None, "Ce fichier n'est pas un modèle KerERC : l'onglet _meta est absent. Téléchargez un modèle à jour.")]
    meta = {str(a): b for a, b in wb["_meta"].iter_rows(max_col=2, values_only=True) if a}
    if str(meta.get("modele") or "") != "kererc-import-projet":
        return {}, [Constat("erreur", None, None, None, "Ce fichier n'est pas un modèle d'import KerERC.")]
    if str(meta.get("version") or "") != VERSION_MODELE:
        return {}, [Constat(
            "erreur", None, None, None,
            "Ce modèle date d'une version antérieure. Téléchargez-en un nouveau.",
        )]
    donnees = {
        "projet": _projet(wb, constats),
        "mesures": _lignes(wb, "Mesures", constats),
        "ug": _lignes(wb, "UG", constats),
        "plan": _lignes(wb, "Actions – plan", constats),
        "calendrier": _lignes(wb, "Actions – calendrier", constats),
        "contacts": _lignes(wb, "Contacts", constats) if "Contacts" in wb.sheetnames else [],
    }
    _valider(donnees, constats)
    return donnees, constats


def lire_octets(contenu: bytes) -> tuple[dict, list[Constat]]:
    from tempfile import NamedTemporaryFile
    with NamedTemporaryFile(suffix=".xlsx") as tmp:
        tmp.write(contenu)
        tmp.flush()
        return lire_classeur(Path(tmp.name))


def _entete(valeur) -> str:
    return texte(valeur).removesuffix("*").strip()


def _projet(wb, constats: list[Constat]) -> dict:
    if "Projet" not in wb.sheetnames:
        constats.append(Constat("erreur", "Projet", None, None, "L'onglet Projet est absent."))
        return {}
    sortie = {}
    for ligne, (cle, valeur) in enumerate(wb["Projet"].iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        nom = _entete(cle)
        if not nom or texte(valeur) == "EXEMPLE" or nom == "EXEMPLE":
            continue
        if texte(valeur).upper() == "EXEMPLE" or str(valeur or "").startswith("EXEMPLE"):
            continue
        sortie[nom] = valeur
    return sortie


def _lignes(wb, onglet: str, constats: list[Constat]) -> list[dict]:
    if onglet not in wb.sheetnames:
        constats.append(Constat("erreur", onglet, None, None, f"L'onglet {onglet} est absent."))
        return []
    feuille = wb[onglet]
    lignes = feuille.iter_rows(values_only=True)
    entete = [_entete(c) for c in next(lignes, [])]
    if not any(entete):
        constats.append(Constat("erreur", onglet, 1, None, "En-têtes absents."))
        return []
    sortie = []
    for numero, valeurs in enumerate(lignes, start=2):
        if not any(v is not None and texte(v) != "" for v in valeurs):
            continue
        if texte(valeurs[0]).upper() == "EXEMPLE" or str(valeurs[0] or "").startswith("EXEMPLE"):
            continue
        ligne = {"_ligne": numero}
        for i, nom in enumerate(entete):
            if not nom:
                continue
            ligne[nom] = valeurs[i] if i < len(valeurs) else None
        sortie.append(ligne)
    return sortie


def _exiger(constats, onglet, ligne, colonne, valeur, message) -> bool:
    if texte(valeur):
        return True
    constats.append(Constat("erreur", onglet, ligne, colonne, message))
    return False


def _valider(donnees: dict, constats: list[Constat]) -> None:
    projet = donnees["projet"]
    for cle in ("Nom", "Commune", "Durée (années)", "Année de début (année N)"):
        if not texte(projet.get(cle)):
            constats.append(Constat("erreur", "Projet", None, cle, f"{cle} est obligatoire."))
    if texte(projet.get("Durée (années)")) and nombre(projet.get("Durée (années)")) is None:
        constats.append(Constat("erreur", "Projet", None, "Durée (années)", "La durée doit être un nombre."))
    if texte(projet.get("Année de début (année N)")) and annee(projet.get("Année de début (année N)")) is None:
        constats.append(Constat("erreur", "Projet", None, "Année de début (année N)", "L'année de début n'est pas une année."))

    mesures = {}
    for ligne in donnees["mesures"]:
        n = ligne["_ligne"]
        c = code(ligne.get("Code"))
        if not _exiger(constats, "Mesures", n, "Code", c, "Le code de la mesure est obligatoire."):
            continue
        if c in mesures:
            constats.append(Constat("erreur", "Mesures", n, "Code", f"Le code {c} est déjà utilisé."))
        mesures[c] = ligne
        _exiger(constats, "Mesures", n, "Libellé", ligne.get("Libellé"), "Le libellé de la mesure est obligatoire.")

    ugs = {}
    for ligne in donnees["ug"]:
        n = ligne["_ligne"]
        c = code(ligne.get("Code"))
        # Le code UG est normalisé plus tard ; l'unicité se juge sur le code saisi compacté.
        if not _exiger(constats, "UG", n, "Code", c, "Le code de l'UG est obligatoire."):
            continue
        if c in ugs:
            constats.append(Constat("erreur", "UG", n, "Code", f"Le code {c} est déjà utilisé."))
        ugs[c] = ligne
        if type_ug(ligne.get("Type")) is None:
            constats.append(Constat("erreur", "UG", n, "Type", "Le type doit être surfacique, linéaire ou ponctuelle."))

    refs = {}
    for ligne in donnees["plan"]:
        n = ligne["_ligne"]
        ref = code(ligne.get("Réf. action"))
        if not _exiger(constats, "Actions – plan", n, "Réf. action", ref, "La référence d'action est obligatoire."):
            continue
        if ref in refs:
            constats.append(Constat("erreur", "Actions – plan", n, "Réf. action", f"{ref} est déjà utilisée."))
        refs[ref] = ligne
        mesure = code(ligne.get("Mesure"))
        if mesure not in mesures:
            constats.append(Constat("erreur", "Actions – plan", n, "Mesure", f"La mesure {mesure or 'vide'} n'existe pas dans l'onglet Mesures."))
        _exiger(constats, "Actions – plan", n, "Libellé", ligne.get("Libellé"), "Le libellé est obligatoire.")
        for ug in liste_codes(ligne.get("UG"), ug=True):
            if not any(code_ug(u.get("Code")) == ug for u in donnees["ug"]):
                constats.append(Constat("erreur", "Actions – plan", n, "UG", f"« {ug} » n'existe pas dans l'onglet UG."))
        if mois(ligne.get("Mois de début")) is None or mois(ligne.get("Mois de fin")) is None:
            constats.append(Constat("erreur", "Actions – plan", n, "Mois de début", "Les mois de début et de fin sont obligatoires (1 à 12 ou le nom du mois)."))
        regle = recurrence(ligne.get("Récurrence"))
        if regle is None:
            constats.append(Constat("erreur", "Actions – plan", n, "Récurrence", "Récurrence inconnue."))
        elif regle == "tous_les_n" and (nombre(ligne.get("N")) or 0) < 1:
            constats.append(Constat("erreur", "Actions – plan", n, "N", "Indiquez N pour « tous les N ans »."))
        elif regle == "annees" and not liste_codes(ligne.get("Années")):
            constats.append(Constat("erreur", "Actions – plan", n, "Années", "Indiquez les années, séparées par un point-virgule."))
        elif regle == "ponctuelle" and annee(ligne.get("Années") or ligne.get("Première année")) is None:
            constats.append(Constat("erreur", "Actions – plan", n, "Années", "Une action ponctuelle demande une année."))
        cout = ligne.get("Coût unitaire HT")
        if texte(cout) and (nombre(cout) is None or nombre(cout) < 0):
            constats.append(Constat("erreur", "Actions – plan", n, "Coût unitaire HT", "Le montant doit être un nombre positif ou nul."))

    for ligne in donnees["calendrier"]:
        n = ligne["_ligne"]
        ref = code(ligne.get("Réf. action"))
        if ref and ref not in refs:
            constats.append(Constat("erreur", "Actions – calendrier", n, "Réf. action", f"{ref} n'existe pas dans l'onglet Actions – plan."))
        if not ref and code(ligne.get("Mesure")) not in mesures:
            constats.append(Constat("erreur", "Actions – calendrier", n, "Mesure", "Une action libre doit citer une mesure de l'onglet Mesures."))
        if not ref and not texte(ligne.get("Libellé")):
            constats.append(Constat("erreur", "Actions – calendrier", n, "Libellé", "Une action libre doit avoir un libellé."))
        if annee(ligne.get("Année")) is None:
            constats.append(Constat("erreur", "Actions – calendrier", n, "Année", "L'année est obligatoire."))
        if statut(ligne.get("Statut")) is None:
            constats.append(Constat("erreur", "Actions – calendrier", n, "Statut", "Statut inconnu."))
        for col in ("Montant prévu HT", "Montant engagé", "Montant réalisé"):
            if texte(ligne.get(col)) and (nombre(ligne.get(col)) is None or nombre(ligne.get(col)) < 0):
                constats.append(Constat("erreur", "Actions – calendrier", n, col, "Le montant doit être un nombre positif ou nul."))

    roles_connus = {lib.lower() for _c, lib in ROLES} | {c for c, _l in ROLES} | {lib.lower() for _c, lib in DISPOSITIFS}
    _ = roles_connus
    for ligne in donnees["contacts"]:
        n = ligne["_ligne"]
        _exiger(constats, "Contacts", n, "Nom", ligne.get("Nom"), "Le nom est obligatoire.")
        _exiger(constats, "Contacts", n, "Rôle", ligne.get("Rôle"), "Le rôle est obligatoire.")


_ = BytesIO
