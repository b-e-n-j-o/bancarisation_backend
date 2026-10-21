"""Lot 1 — formats, nommage, absence de document, formules budget. Sans base."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from api.calendrier_ics import generer_ics
from api.export.archive import ecrire_zip
from api.export.contexte import ExportContexte, ExportOptions, FichierExport
from api.export.manifest import construire, ecrire, empreinte
from api.export.modules.budget import ecrire_budget_xlsx, produire as produire_budget
from api.export.modules.documents import DocumentAbsent, LigneDocument, deposer_documents
from api.export.nommage import racine_zip, reserver, slug
from api.export.sqlutil import pour_json


def _ctx(tmp: Path, **options) -> ExportContexte:
    return ExportContexte(
        export_id=__import__("uuid").uuid4(),
        projet_id=__import__("uuid").uuid4(),
        user_id=__import__("uuid").uuid4(),
        organisation_id=__import__("uuid").uuid4(),
        options=ExportOptions(**options),
        dossier_tmp=tmp,
        url_front="https://bancarisation.kerelia.fr",
        racine_zip=racine_zip("REF-01", "abcdef1234567890", "Étang des Îles", "2026-10-21"),
    )


def test_slug_ascii():
    assert slug("Étang des Îles") == "etang-des-iles"
    assert slug("") == "sans-nom"


def test_chemins_courts_et_ascii():
    pris: set[str] = set()
    racine = racine_zip(None, "abcdef1234567890", "Nom", "2026-10-21")
    long = "02_Calendrier/" + ("a" * 200) + ".xlsx"
    chemin = reserver(long, racine, pris)
    assert len(f"{racine}/{chemin}") <= 180
    assert chemin.isascii()
    doublon = reserver("02_Calendrier/actions.xlsx", racine, pris)
    autre = reserver("02_Calendrier/actions.xlsx", racine, pris)
    assert doublon != autre
    assert autre.endswith("_2.xlsx") or "_2" in Path(autre).stem


def test_ics_fenetre_longue_deux_evenements():
    brut = generer_ics(
        [{
            "id": "occ-1",
            "annee": 2026,
            "mois_debut": 1,
            "mois_fin": 4,
            "statut": "a_confirmer",
            "titre": "Fauchage",
            "mesure": "C1 Prairie",
            "ug": "UG1",
            "montant_ht": 1500,
            "prestataire": "Secret",
        }],
        nom_projet="Étang",
        url_front="https://bancarisation.kerelia.fr",
        projet_id="p1",
        mode_titre="complet",
    )
    from icalendar import Calendar

    cal = Calendar.from_ical(brut)
    events = [c for c in cal.walk("VEVENT")]
    assert len(events) == 2
    uids = {str(e.get("uid")) for e in events}
    assert "occ-1-ouverture@kererc.fr" in uids
    assert "occ-1-echeance@kererc.fr" in uids
    assert str(events[0].get("status")) == "TENTATIVE"
    description = str(events[0].get("description"))
    assert "montant" not in description.lower()
    assert "Secret" not in description
    assert "vue=calendar" in description
    assert "occurrence=occ-1" in description


def test_ics_fenetre_courte_un_evenement():
    brut = generer_ics(
        [{
            "id": "occ-2",
            "annee": 2026,
            "mois_debut": 3,
            "mois_fin": 3,
            "statut": "planifie",
            "titre": "Visite",
        }],
        nom_projet="Site",
        url_front="https://bancarisation.kerelia.fr",
        projet_id="p1",
    )
    from icalendar import Calendar

    events = [c for c in Calendar.from_ical(brut).walk("VEVENT")]
    assert len(events) == 1
    assert str(events[0].get("uid")) == "occ-2@kererc.fr"
    assert events[0].get("dtstart").dt == date(2026, 3, 1)


def test_vcard():
    from api.export.modules import contacts
    texte = contacts._vcard(
        [{
            "nom": "Durand",
            "prenom": "Léa",
            "structure": "Kerelia",
            "fonction": "Chargée",
            "email": "lea@exemple.fr",
            "telephone": "0600000000",
            "role": "Maître d'ouvrage",
        }],
        "Étang",
    )
    import vobject

    cartes = list(vobject.readComponents(texte))
    assert len(cartes) == 1
    assert cartes[0].version.value == "3.0"
    assert "Durand" in cartes[0].fn.value


def test_budget_formules(tmp_path: Path):
    ecrire_budget_xlsx(
        tmp_path / "budget.xlsx",
        {
            "detail": [{
                "id": "1",
                "mesure_code": "C1",
                "mesure_libelle": "Prairie",
                "titre": "Fauchage",
                "ug_ids": ["UG1"],
                "annee": 2026,
                "montant_ht": 100,
                "montant_engage": 40,
                "montant_realise": 10,
                "montant_initial": 80,
                "annee_initiale": 2026,
                "prestataire": "Atelier",
                "statut": "planifie",
            }],
            "baselines": [],
            "ecarts": [],
            "mouvements": [],
        },
    )
    wb = load_workbook(tmp_path / "budget.xlsx")
    synthese = wb["Synthèse annuelle"]
    assert str(synthese["B2"].value).startswith("=SUMIFS(")
    assert wb.calculation.fullCalcOnLoad is True


def test_budget_exclu(tmp_path: Path):
    ctx = _ctx(tmp_path, inclure_budget=False)
    assert produire_budget(ctx, None) == []  # type: ignore[arg-type]


def test_document_manquant_avertit_sans_lever(tmp_path: Path):
    ctx = _ctx(tmp_path, inclure_documents=True)

    def telecharger(_chemin: str, _cible: Path) -> str:
        raise DocumentAbsent("objet absent")

    produits = deposer_documents(
        ctx,
        [LigneDocument(
            id="d1",
            nom="Arrêté.pdf",
            nom_fichier="Arrete.pdf",
            bucket_path="proj/arrete.pdf",
            categorie="arrete",
            date_document=date(2026, 1, 2),
            created_at=None,
            taille_octets=10,
            deposant=None,
            occurrence_id=None,
            liens="",
        )],
        telecharger,
        True,
    )
    assert ctx.avertissements
    assert any(p.chemin_archive.endswith("index_documents.xlsx") for p in produits)
    wb = load_workbook(next(p.chemin_local for p in produits if p.chemin_archive.endswith(".xlsx")))
    valeurs = [cellule.value for ligne in wb.active.iter_rows(min_row=2) for cellule in ligne]
    assert "fichier manquant" in valeurs


def test_pyogrio_gpkg(tmp_path: Path):
    pyogrio = __import__("pytest").importorskip("pyogrio")
    assert "GPKG" in pyogrio.list_drivers()
    from shapely.geometry import Point
    from shapely import wkb as shapely_wkb

    from api.export.modules.geometries import _ecrire_gpkg

    brut = shapely_wkb.dumps(Point(500000, 6500000))
    cible = tmp_path / "projet.gpkg"
    _ecrire_gpkg(cible, "emprise_projet", [{"_wkb": brut, "nom": "Emprise"}], ["nom"])
    couches = [ligne[0] for ligne in pyogrio.list_layers(cible)]
    assert "emprise_projet" in couches


def test_manifest_et_zip(tmp_path: Path):
    ctx = _ctx(tmp_path)
    source = tmp_path / "note.txt"
    source.write_text("bonjour", encoding="utf-8")
    fic = FichierExport(
        chemin_archive="01_Projet/note.txt",
        chemin_local=source,
        module="projet",
        description="Note",
        nb_elements=1,
    )
    manifest = construire(
        ctx,
        [fic],
        {"projet": {"statut": "ok"}},
        email="a@exemple.fr",
        organisation_nom="Org",
        projet_nom="Étang",
        reference="REF-01",
    )
    assert "manifest.json" not in {e["chemin"] for e in manifest["fichiers"]}
    chemin_manifest = ecrire(tmp_path, manifest)
    sha, _taille = empreinte(source)
    assert manifest["fichiers"][0]["sha256"] == sha
    zip_path = ecrire_zip(tmp_path, ctx.racine_zip, [fic], chemin_manifest)
    import zipfile

    with zipfile.ZipFile(zip_path) as archive:
        noms = archive.namelist()
    assert any(n.endswith("manifest.json") for n in noms)
    assert any(n.endswith("note.txt") for n in noms)
    assert all(Path(n).name.isascii() and n.isascii() for n in noms)
    assert all(len(n) <= 180 for n in noms)
    assert pour_json({"a": 1}) == {"a": 1}
