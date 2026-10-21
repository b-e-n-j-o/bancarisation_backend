"""03_Budget : classeur à formules et historique CSV."""

from __future__ import annotations

import csv
from pathlib import Path

import psycopg
from openpyxl.utils import get_column_letter

from api.budget.lecture_export import charger_budget_projet
from api.export.contexte import ExportContexte, FichierExport
from api.export.excel_util import ecrire_feuille, libelle_statut, nouveau_classeur, sauver
from api.export.modules._fichiers import fichier
from api.export.sqlutil import pour_json

nom = "budget"
dossier = "03_Budget"
_FEUILLE_DETAIL = "Détail par action"


def produire(ctx: ExportContexte, conn: psycopg.Connection) -> list[FichierExport]:
    if not ctx.options.inclure_budget:
        return []
    data = charger_budget_projet(
        conn,
        str(ctx.projet_id),
        ctx.options.annee_min,
        ctx.options.annee_max,
    )
    fic = fichier(ctx, f"{dossier}/budget.xlsx", nom, "Budget détaillé et synthèses", len(data["detail"]))
    ecrire_budget_xlsx(fic.chemin_local, data)
    historique = fichier(
        ctx,
        f"{dossier}/historique_mouvements.csv",
        nom,
        "Historique des mouvements budgétaires",
        len(data["mouvements"]),
    )
    _ecrire_csv(historique.chemin_local, data["mouvements"])
    return [fic, historique]


def ecrire_budget_xlsx(chemin: Path, data: dict) -> None:
    wb = nouveau_classeur()
    detail = data["detail"]
    ws = wb.active
    ws.title = _FEUILLE_DETAIL
    lignes_detail = [
        [
            r.get("id"),
            " ".join(p for p in (r.get("mesure_code"), r.get("mesure_libelle")) if p).strip(),
            r.get("titre"),
            ", ".join(r.get("ug_ids") or []),
            r.get("annee"),
            _nombre(r.get("montant_ht")),
            _nombre(r.get("montant_engage")),
            _nombre(r.get("montant_realise")),
            _nombre(r.get("montant_initial")),
            r.get("annee_initiale"),
            r.get("prestataire") or "",
            libelle_statut(r.get("statut")),
        ]
        for r in detail
    ]
    ecrire_feuille(
        ws,
        [
            "Id", "Mesure", "Action", "UG", "Année",
            "Prévu HT", "Engagé", "Réalisé", "Initial", "Année initiale",
            "Prestataire", "Statut",
        ],
        lignes_detail,
        colonnes_montant=(6, 7, 8, 9),
        colonne_statut=12,
    )

    annees = sorted({r.get("annee") for r in detail if r.get("annee") is not None})
    synthese = wb.create_sheet("Synthèse annuelle")
    synthese.append(["Année", "Prévu", "Engagé", "Réalisé", "Initial", "Écart prévu / initial"])
    for i, annee in enumerate(annees, start=2):
        synthese.cell(i, 1, annee)
        for col, source in ((2, "F"), (3, "G"), (4, "H")):
            synthese.cell(
                i, col,
                f"=SUMIFS('{_FEUILLE_DETAIL}'!{source}:{source},'{_FEUILLE_DETAIL}'!E:E,A{i})",
            )
        synthese.cell(
            i, 5,
            f"=SUMIFS('{_FEUILLE_DETAIL}'!I:I,'{_FEUILLE_DETAIL}'!J:J,A{i})",
        )
        synthese.cell(i, 6, f"=B{i}-E{i}")
    total_row = len(annees) + 2
    synthese.cell(total_row, 1, "Total")
    for col in range(2, 7):
        lettre = get_column_letter(col)
        synthese.cell(total_row, col, f"=SUM({lettre}2:{lettre}{total_row - 1})")
    _habiller_entete(synthese)

    prestataires = wb.create_sheet("Par prestataire")
    prestataires.append(["Prestataire", "Statut", "Prévu", "Engagé", "Réalisé"])
    couples = sorted({
        ((r.get("prestataire") or "Non renseigné"), libelle_statut(r.get("statut")))
        for r in detail
    })
    for i, (prestataire, statut) in enumerate(couples, start=2):
        prestataires.cell(i, 1, prestataire)
        prestataires.cell(i, 2, statut)
        for col, source in ((3, "F"), (4, "G"), (5, "H")):
            prestataires.cell(
                i, col,
                (
                    f"=SUMIFS('{_FEUILLE_DETAIL}'!{source}:{source},"
                    f"'{_FEUILLE_DETAIL}'!K:K,A{i},'{_FEUILLE_DETAIL}'!L:L,B{i})"
                ),
            )
    _habiller_entete(prestataires)

    base = wb.create_sheet("Baseline")
    ecrire_feuille(
        base,
        ["Id", "Figé le", "Libellé", "Commentaire", "Mode", "Occurrences", "Total HT"],
        [
            [
                b.get("id"), b.get("figee_le"), b.get("libelle"), b.get("commentaire"),
                b.get("mode"), b.get("nb_occurrences"), _nombre(b.get("total_ht")),
            ]
            for b in data["baselines"]
        ],
        colonnes_date=(2,),
        colonnes_montant=(7,),
    )

    ecarts = wb.create_sheet("Écarts")
    ecrire_feuille(
        ecarts,
        ["Année", "Annulé", "Glissé sortant", "Ajouté", "Révision de prix"],
        [
            [
                e.get("annee_ref"),
                _nombre(e.get("annule")),
                _nombre(e.get("glisse_sortant")),
                _nombre(e.get("ajoute")),
                _nombre(e.get("revision_prix")),
            ]
            for e in data["ecarts"]
        ],
        colonnes_montant=(2, 3, 4, 5),
    )

    histo = wb.create_sheet("Historique")
    ecrire_feuille(
        histo,
        ["Date", "Occurrence", "Champ", "Ancienne valeur", "Nouvelle valeur", "Motif", "Auteur"],
        [
            [
                m.get("modifie_le"), m.get("occurrence"), m.get("champ"),
                m.get("ancienne_val"), m.get("nouvelle_val"), m.get("motif"), m.get("modifie_par"),
            ]
            for m in data["mouvements"]
        ],
        colonnes_date=(1,),
    )
    sauver(wb, chemin)


def _habiller_entete(ws) -> None:
    from openpyxl.styles import Font, PatternFill

    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="245C3A")
    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions


def _nombre(valeur):
    if valeur is None:
        return None
    return float(valeur)


def _ecrire_csv(chemin: Path, mouvements: list[dict]) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    champs = ["modifie_le", "occurrence_id", "occurrence", "champ", "ancienne_val", "nouvelle_val", "motif", "modifie_par"]
    with chemin.open("w", encoding="utf-8", newline="") as fichier_csv:
        writer = csv.DictWriter(fichier_csv, fieldnames=champs, extrasaction="ignore")
        writer.writeheader()
        for ligne in mouvements:
            writer.writerow({cle: pour_json(ligne.get(cle)) for cle in champs})
