"""Classeur Excel lisible : en-têtes, filtre, dates, montants, statuts."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

STATUTS = {
    "a_confirmer": "À confirmer",
    "planifie": "Planifiée",
    "en_cours": "En cours",
    "realise": "Réalisée",
    "repousse": "Repoussée",
    "archive": "Archivée",
    "supprime": "Supprimée",
    "brouillon": "Brouillon",
    "valide": "Validé",
    "identifie": "Identifié",
    "en_negociation": "En négociation",
    "signe": "Signé",
    "publie": "Publié",
    "expire": "Expiré",
    "resilie": "Résilié",
}

_COULEURS_STATUT = {
    "À confirmer": "FFF4E4",
    "Planifiée": "E7F0E4",
    "En cours": "E4EEF8",
    "Réalisée": "D9F2E3",
    "Repoussée": "F8E8D4",
    "Archivée": "EEEEEE",
    "Supprimée": "F8D6D6",
}

_ENTETE = Font(bold=True, color="FFFFFF")
_FOND_ENTETE = PatternFill("solid", fgColor="245C3A")
_FORMAT_DATE = "DD/MM/YYYY"
_FORMAT_MONTANT = '#,##0.00 "€"'


def nouveau_classeur() -> Workbook:
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    return wb


def ecrire_feuille(
    ws: Worksheet,
    entetes: Sequence[str],
    lignes: Sequence[Sequence[Any]],
    *,
    colonnes_date: Sequence[int] = (),
    colonnes_montant: Sequence[int] = (),
    colonne_statut: int | None = None,
) -> None:
    ws.append(list(entetes))
    for ligne in lignes:
        ws.append([_cellule(v) for v in ligne])
    for cellule in ws[1]:
        cellule.font = _ENTETE
        cellule.fill = _FOND_ENTETE
        cellule.alignment = Alignment(vertical="center")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}{max(ws.max_row, 1)}"
    for index in colonnes_date:
        for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            if isinstance(row[0].value, (date, datetime)):
                row[0].number_format = _FORMAT_DATE
    for index in colonnes_montant:
        for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            if isinstance(row[0].value, (int, float, Decimal)):
                row[0].number_format = _FORMAT_MONTANT
    if colonne_statut is not None:
        _colorer_statut(ws, colonne_statut)
    _largeurs(ws)


def libelle_statut(valeur: Any) -> str:
    if valeur is None:
        return ""
    return STATUTS.get(str(valeur), str(valeur).replace("_", " ").capitalize())


def sauver(wb: Workbook, chemin: Path) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)


def _cellule(valeur: Any) -> Any:
    if isinstance(valeur, Decimal):
        return float(valeur)
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, (list, dict)):
        return ", ".join(str(v) for v in valeur) if isinstance(valeur, list) else str(valeur)
    return valeur


def _colorer_statut(ws: Worksheet, colonne: int) -> None:
    from openpyxl.formatting.rule import FormulaRule

    lettre = get_column_letter(colonne)
    for libelle, couleur in _COULEURS_STATUT.items():
        formule = f'{lettre}2="{libelle}"'
        ws.conditional_formatting.add(
            f"{lettre}2:{lettre}{max(ws.max_row, 2)}",
            FormulaRule(
                formula=[formule],
                fill=PatternFill("solid", fgColor=couleur),
            ),
        )


def _largeurs(ws: Worksheet) -> None:
    for index, colonne in enumerate(ws.columns, start=1):
        largeur = 12
        for cellule in colonne[:80]:
            if cellule.value is not None:
                largeur = max(largeur, min(len(str(cellule.value)) + 2, 42))
        ws.column_dimensions[get_column_letter(index)].width = largeur
